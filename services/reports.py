"""
Service de génération de rapports PDF et Excel.
Utilise ReportLab pour PDF et openpyxl pour Excel.
"""

import os
import logging
from datetime import datetime, timezone
from typing import Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from models.orm import InspectionSession, InspectionResult, Lot, Product, Report, DecisionEnum
from config import settings

logger = logging.getLogger(__name__)

os.makedirs(settings.EXPORTS_DIR, exist_ok=True)


# ── PDF ───────────────────────────────────────────────────────────────────

async def generer_rapport_pdf(
    db: AsyncSession,
    session_id: Optional[int] = None,
    lot_id: Optional[int] = None,
    report_type: str = "session"
) -> str:
    """
    Génère un rapport PDF et retourne le chemin du fichier.
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph,
            Spacer, HRFlowable
        )
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
    except ImportError:
        raise RuntimeError("ReportLab non installé. Lancez : pip install reportlab")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"rapport_{report_type}_{timestamp}.pdf"
    filepath = os.path.join(settings.EXPORTS_DIR, filename)

    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4,
        rightMargin=2*cm,
        leftMargin=2*cm,
        topMargin=2*cm,
        bottomMargin=2*cm
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title", parent=styles["Heading1"],
        fontSize=18, textColor=colors.HexColor("#0F1117"),
        spaceAfter=6, alignment=TA_CENTER
    )
    subtitle_style = ParagraphStyle(
        "Subtitle", parent=styles["Normal"],
        fontSize=10, textColor=colors.grey,
        spaceAfter=12, alignment=TA_CENTER
    )
    heading_style = ParagraphStyle(
        "Heading", parent=styles["Heading2"],
        fontSize=13, textColor=colors.HexColor("#00C2FF"),
        spaceBefore=12, spaceAfter=6
    )
    normal_style = styles["Normal"]

    story = []

    # ── En-tête ──
    story.append(Paragraph(settings.COMPANY_NAME, title_style))
    story.append(Paragraph("PharmaVision AI — Rapport d'Inspection", subtitle_style))
    story.append(Paragraph(
        f"Généré le : {datetime.now().strftime('%d/%m/%Y à %H:%M:%S')}",
        subtitle_style
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#00C2FF")))
    story.append(Spacer(1, 0.5*cm))

    # ── Données selon le type ──
    if report_type == "session" and session_id:
        session = await db.get(InspectionSession, session_id)
        if session:
            lot = await db.get(Lot, session.lot_id)
            product = await db.get(Product, lot.product_id) if lot else None

            story.append(Paragraph("Informations de la Session", heading_style))
            info_data = [
                ["Session ID", str(session.id)],
                ["Produit", product.name if product else "N/A"],
                ["Numéro de lot", lot.lot_number if lot else "N/A"],
                ["Démarré le", session.started_at.strftime("%d/%m/%Y %H:%M:%S") if session.started_at else "N/A"],
                ["Arrêté le", session.stopped_at.strftime("%d/%m/%Y %H:%M:%S") if session.stopped_at else "En cours"],
                ["Total inspecté", str(session.total_inspected)],
                ["Conformes", str(session.total_conforme)],
                ["Suspects", str(session.total_suspect)],
                ["Non conformes", str(session.total_non_conforme)],
                ["Taux de conformité", f"{session.conformity_rate:.2f}%"],
            ]
            info_table = Table(info_data, colWidths=[5*cm, 12*cm])
            info_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#1E2130")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (1, 0), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(info_table)
            story.append(Spacer(1, 0.5*cm))

            # Résultats individuels
            story.append(Paragraph("Résultats Détaillés", heading_style))
            result_query = select(InspectionResult).where(
                InspectionResult.session_id == session_id
            ).order_by(InspectionResult.timestamp)
            result_rows = await db.execute(result_query)
            results = result_rows.scalars().all()

            if results:
                table_data = [["#", "Heure", "Décision", "Score Anomalie", "Confiance YOLO", "Durée (ms)"]]
                for i, r in enumerate(results, 1):
                    decision_label = r.decision.value if r.decision else "N/A"
                    table_data.append([
                        str(i),
                        r.timestamp.strftime("%H:%M:%S") if r.timestamp else "N/A",
                        decision_label,
                        f"{r.anomaly_score:.4f}" if r.anomaly_score is not None else "N/A",
                        f"{r.yolo_confidence:.4f}" if r.yolo_confidence is not None else "N/A",
                        f"{r.analysis_duration_ms:.1f}" if r.analysis_duration_ms is not None else "N/A",
                    ])

                results_table = Table(table_data, colWidths=[1*cm, 2.5*cm, 3.5*cm, 3.5*cm, 3.5*cm, 3*cm])
                results_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00C2FF")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(results_table)

    # ── Pied de page ──
    story.append(Spacer(1, 1*cm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
    story.append(Paragraph(
        f"Document généré par {settings.APP_NAME} v{settings.APP_VERSION} — Confidentiel",
        ParagraphStyle("Footer", parent=normal_style, fontSize=7, textColor=colors.grey, alignment=TA_CENTER)
    ))

    doc.build(story)
    logger.info(f"Rapport PDF généré : {filepath}")
    return filepath


# ── Excel ─────────────────────────────────────────────────────────────────

async def generer_rapport_excel(
    db: AsyncSession,
    session_id: Optional[int] = None,
    lot_id: Optional[int] = None,
    report_type: str = "session"
) -> str:
    """
    Génère un rapport Excel et retourne le chemin du fichier.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl non installé. Lancez : pip install openpyxl")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"rapport_{report_type}_{timestamp}.xlsx"
    filepath = os.path.join(settings.EXPORTS_DIR, filename)

    wb = openpyxl.Workbook()

    # Couleurs
    BLUE = "00C2FF"
    DARK = "0F1117"
    GREEN = "22C55E"
    ORANGE = "F59E0B"
    RED = "EF4444"
    LIGHT_GRAY = "F8FAFC"

    def header_style(cell, bg=BLUE):
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def decision_color(decision: str) -> str:
        return {"CONFORME": GREEN, "SUSPECT": ORANGE, "NON_CONFORME": RED}.get(decision, "CCCCCC")

    # ── Feuille Résumé ──
    ws_summary = wb.active
    ws_summary.title = "Résumé"

    ws_summary["A1"] = settings.COMPANY_NAME
    ws_summary["A1"].font = Font(bold=True, size=16, color=DARK)
    ws_summary["A2"] = "PharmaVision AI — Rapport d'Inspection"
    ws_summary["A3"] = f"Généré le : {datetime.now().strftime('%d/%m/%Y à %H:%M:%S')}"
    ws_summary["A3"].font = Font(italic=True, color="888888", size=9)

    if report_type == "session" and session_id:
        session = await db.get(InspectionSession, session_id)
        if session:
            lot = await db.get(Lot, session.lot_id)
            product = await db.get(Product, lot.product_id) if lot else None

            ws_summary["A5"] = "INFORMATIONS SESSION"
            header_style(ws_summary["A5"], DARK)

            infos = [
                ("Session ID", str(session.id)),
                ("Produit", product.name if product else "N/A"),
                ("Numéro de lot", lot.lot_number if lot else "N/A"),
                ("Démarré le", session.started_at.strftime("%d/%m/%Y %H:%M:%S") if session.started_at else "N/A"),
                ("Arrêté le", session.stopped_at.strftime("%d/%m/%Y %H:%M:%S") if session.stopped_at else "En cours"),
                ("Total inspecté", session.total_inspected),
                ("Conformes", session.total_conforme),
                ("Suspects", session.total_suspect),
                ("Non conformes", session.total_non_conforme),
                ("Taux de conformité", f"{session.conformity_rate:.2f}%"),
            ]
            for i, (key, val) in enumerate(infos, 6):
                ws_summary[f"A{i}"] = key
                ws_summary[f"A{i}"].font = Font(bold=True)
                ws_summary[f"B{i}"] = val
                if i % 2 == 0:
                    ws_summary[f"A{i}"].fill = PatternFill("solid", fgColor=LIGHT_GRAY)
                    ws_summary[f"B{i}"].fill = PatternFill("solid", fgColor=LIGHT_GRAY)

            ws_summary.column_dimensions["A"].width = 25
            ws_summary.column_dimensions["B"].width = 30

            # ── Feuille Résultats détaillés ──
            ws_results = wb.create_sheet("Résultats Détaillés")
            headers = ["#", "Date/Heure", "Décision", "Score Anomalie",
                       "Confiance YOLO", "Durée (ms)", "BBox X1", "BBox Y1", "BBox X2", "BBox Y2"]
            for col, h in enumerate(headers, 1):
                cell = ws_results.cell(row=1, column=col, value=h)
                header_style(cell)
                ws_results.column_dimensions[get_column_letter(col)].width = 16

            result_query = select(InspectionResult).where(
                InspectionResult.session_id == session_id
            ).order_by(InspectionResult.timestamp)
            result_rows = await db.execute(result_query)
            results = result_rows.scalars().all()

            for row_idx, r in enumerate(results, 2):
                decision_val = r.decision.value if r.decision else "N/A"
                row_data = [
                    row_idx - 1,
                    r.timestamp.strftime("%d/%m/%Y %H:%M:%S") if r.timestamp else "N/A",
                    decision_val,
                    round(r.anomaly_score, 4) if r.anomaly_score is not None else None,
                    round(r.yolo_confidence, 4) if r.yolo_confidence is not None else None,
                    round(r.analysis_duration_ms, 1) if r.analysis_duration_ms is not None else None,
                    r.bbox_x1, r.bbox_y1, r.bbox_x2, r.bbox_y2
                ]
                for col_idx, val in enumerate(row_data, 1):
                    cell = ws_results.cell(row=row_idx, column=col_idx, value=val)
                    cell.alignment = Alignment(horizontal="center")
                    if col_idx == 3:  # Décision
                        dc = decision_color(decision_val)
                        cell.fill = PatternFill("solid", fgColor=dc)
                        cell.font = Font(color="FFFFFF", bold=True)
                    if row_idx % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)

    wb.save(filepath)
    logger.info(f"Rapport Excel généré : {filepath}")
    return filepath


async def exporter_inspections_csv(
    db: AsyncSession,
    session_id: Optional[int] = None,
    lot_id: Optional[int] = None
) -> str:
    """Export CSV simple pour intégration ERP/LIMS."""
    import csv

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"export_inspections_{timestamp}.csv"
    filepath = os.path.join(settings.EXPORTS_DIR, filename)

    query = select(InspectionResult)
    if session_id:
        query = query.where(InspectionResult.session_id == session_id)

    rows = await db.execute(query)
    results = rows.scalars().all()

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "id", "session_id", "timestamp", "decision",
            "anomaly_score", "yolo_confidence", "analysis_duration_ms",
            "bbox_x1", "bbox_y1", "bbox_x2", "bbox_y2"
        ])
        for r in results:
            writer.writerow([
                r.id, r.session_id,
                r.timestamp.isoformat() if r.timestamp else "",
                r.decision.value if r.decision else "",
                r.anomaly_score, r.yolo_confidence, r.analysis_duration_ms,
                r.bbox_x1, r.bbox_y1, r.bbox_x2, r.bbox_y2
            ])

    logger.info(f"Export CSV généré : {filepath} ({len(results)} lignes)")
    return filepath
